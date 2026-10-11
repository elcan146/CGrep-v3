"""Compare sdrplay_real.py CSV output against an SDRuno export.

Usage:
    python compare_sdruno.py <script_csv> <sdruno_file>

- script_csv : sdrplay_real_*.csv produced by sdrplay_real.py
               (columns: center_freq_hz, peak_dbm, floor_dbm, snr_db)
- sdruno_file: SDRuno export — .xlsx with columns Tarix/Tezlik/Signal/SNR
               (Tezlik in Hz) OR .csv/.txt with freq_hz,signal,snr
               (tab/comma/space separated, header optional)

Prints per-frequency differences (mine - SDRuno) and summary stats.
"""
import sys
import csv
import numpy as np


def load_script_csv(path):
    """freq_kHz -> list of (pk, floor, snr)"""
    data = {}
    with open(path, newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            try:
                khz = int(round(float(row["center_freq_hz"]) / 1000.0))
                data.setdefault(khz, []).append(
                    (float(row["peak_dbm"]), float(row["floor_dbm"]), float(row["snr_db"]))
                )
            except (KeyError, ValueError):
                continue
    return data


def load_sdruno(path):
    """freq_kHz -> list of (signal, snr)"""
    data = {}
    if path.lower().endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = [str(c).strip().lower() if c else "" for c in next(rows)]
        # find columns: Tezlik (freq Hz), Signal, SNR
        try:
            i_f = header.index("tezlik")
            i_s = header.index("signal")
            i_n = header.index("snr")
        except ValueError:
            i_f, i_s, i_n = 1, 2, 3
        for r in rows:
            try:
                khz = int(round(float(r[i_f]) / 1000.0))
                data.setdefault(khz, []).append((float(r[i_s]), float(r[i_n])))
            except (TypeError, ValueError, IndexError):
                continue
    else:
        with open(path, encoding="utf-8", errors="replace") as f:
            for line in f:
                parts = line.replace(",", " ").split()
                if len(parts) < 3:
                    continue
                try:
                    freq = float(parts[0])
                    sig = float(parts[1])
                    snr = float(parts[-1])
                except ValueError:
                    continue
                khz = int(round(freq / 1000.0)) if freq > 10000 else int(round(freq))
                data.setdefault(khz, []).append((sig, snr))
    return data


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return
    mine = load_script_csv(sys.argv[1])
    ref = load_sdruno(sys.argv[2])

    common = sorted(set(mine) & set(ref))
    if not common:
        print("No overlapping frequencies found.")
        return

    print(f"{'MHz':>8} | {'Pk mine':>8} {'Pk SDRuno':>9} {'dPk':>6} | "
          f"{'SNR mine':>8} {'SNR SDRuno':>9} {'dSNR':>6}")
    print("-" * 70)
    dpk, dsnr = [], []
    for khz in common:
        m = np.mean(mine[khz], axis=0)
        r = np.mean(ref[khz], axis=0)
        dp = m[0] - r[0]
        ds = m[2] - r[1]
        dpk.append(dp)
        dsnr.append(ds)
        flag = "  <-- BIG" if abs(dp) > 5 else ""
        print(f"{khz/1000:8.1f} | {m[0]:8.2f} {r[0]:9.2f} {dp:+6.2f} | "
              f"{m[2]:8.2f} {r[1]:9.2f} {ds:+6.2f}{flag}")

    dpk, dsnr = np.array(dpk), np.array(dsnr)
    print("-" * 70)
    print(f"Freqs compared: {len(common)}")
    print(f"dPk : mean {dpk.mean():+.2f} dB, mean|d| {np.abs(dpk).mean():.2f} dB, "
          f"max|d| {np.abs(dpk).max():.2f} dB")
    print(f"dSNR: mean {dsnr.mean():+.2f} dB, mean|d| {np.abs(dsnr).mean():.2f} dB, "
          f"max|d| {np.abs(dsnr).max():.2f} dB")


if __name__ == "__main__":
    main()
